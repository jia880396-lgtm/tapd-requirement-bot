# -*- coding: utf-8 -*-
"""生成产品经理简历「实习经历」段落 Word 文档（纯黑文字，精简版）。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BLACK = RGBColor(0, 0, 0)
doc = Document()

style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(11)
style.font.color.rgb = BLACK
import docx.oxml.ns as ns
style.element.rPr.rFonts.set(ns.qn("w:eastAsia"), "宋体")

def black(run):
    run.font.color.rgb = BLACK

def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    black(r)
    return p

def sub(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    black(r)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(11)
    black(r)
    return p

# 标题
t = doc.add_heading("实习经历（可直接粘贴至简历）", level=0)
for r in t.runs:
    black(r)

# 经历块头
h1("旺店通（电商 ERP SaaS）  |  产品经理实习生  |  2026.06 – 2026.08（请按实际调整）")

# 项目一句话背景（精简）
sub("项目简介：主导搭建一套面向旺店通 ERP 的 AI 需求处理工具，覆盖需求分类、可靠性评估、"
    "重复识别、PRD 评审四大环节，已落地至团队 20+ 产品经理的日常需求评审流程。")

# 五大维度（精简、聚焦核心）
bullet("业务梳理与知识库建设：基于 2.2 万条历史需求统计，搭建覆盖 22 个业务模块、40+ 专业术语的"
       "结构化业务知识库，为 AI 判断提供领域依据。")

bullet("评估体系与提示词设计：独立设计需求可靠性四维评分体系（信息完整度/描述清晰度/可实现性/业务价值）"
       "与宽松-标准-严格三档策略，将模糊标准改写为「事实核查清单+档位锚点」判定规则，模型打分与证据自洽性显著提升。")

bullet("模型选型与成本优化：主导图片理解方案选型，实测对比多家模型供应商（价差最高 300 倍），"
       "推动「Qwen3-VL 一体化」方案落地——中文 OCR 准确率 87%→96%，单条需求处理耗时降低 16%。")

bullet("迭代机制与数据闭环：搭建提示词版本化管理 + A/B 评测 + 一键回滚机制，通过 10 组双方案对照实验"
       "持续迭代，打分漂移收敛 50%；设计「人工打分→差异案例自动沉淀→提示词再优化」自迭代闭环。")

bullet("产品落地与跨团队协作：经小范围验证后推广至 20+ 产品经理日常流程，降低需求评审工作量、"
       "加速客户需求分派至对应负责人、提升 PRD 撰写质量，并持续迭代优化。")

# 工具与方法（可选）
h1("工具与方法（可选，置于简历技能区）")
bullet("需求调研与 PRD：用户访谈、竞品分析、PRD 撰写、Axure/Figma 原型")
bullet("数据分析与实验：埋点设计、SQL、指标体系搭建、A/B 测试")
bullet("AI 产品化：大模型 Prompt 设计、案例库与评测集构建、模型效果度量与回归管控")

out = r"../docs/产品经理实习经历_旺店通.docx"
doc.save(out)
print("saved:", out)
